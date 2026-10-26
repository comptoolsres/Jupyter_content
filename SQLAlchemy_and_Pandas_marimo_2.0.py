import marimo

__generated_with = "0.24.0"
app = marimo.App()


@app.cell
def _():
    import marimo as mo
    import pandas as pd
    from openpyxl import load_workbook
    from sqlalchemy import MetaData, Table, create_engine, select

    return MetaData, Table, create_engine, load_workbook, mo, pd, select


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    # Pandas and SQLAlchemy (2.0 style)

    Pandas has several functions to work with databases via SQLAlchemy.

    This can be an easier way to get data from a database and to pass data into a database.

    This notebook uses SQLAlchemy 2.x Core patterns and requires pandas 2.x or later.

    Let's start with loading the `world.sqlite` database we've been using.
    """)
    return


@app.cell
def _(MetaData, create_engine):
    # Create an Engine object which is our handle into the database.
    engine = create_engine('sqlite:////blue/bsc4452/share/Class_Files/data/world.sqlite')

    # MetaData holds table definitions reflected from the existing database.
    metadata = MetaData()
    return engine, metadata


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The cell below is also from the [SQLAlchemy 2.0 notebook](SQLAlchemy_2.0.ipynb). It creates the `city` object by reflecting the table from the database (using `autoload_with=engine`) and builds the `query` variable with `select()`. It then shows how you would execute the query and print the rows, using a connection in a `with` block so it is closed automatically.
    """)
    return


@app.cell
def _(Table, engine, metadata, select):
    # Auto-create the city object by reflecting the table from the database.
    city = Table('city', metadata, autoload_with=engine)

    # Mirroring this command: SELECT Id, Name, Population FROM city LIMIT 10;
    query = select(city.c.ID, city.c.Name, city.c.Population).limit(10)

    with engine.connect() as _conn:
        for _row in _conn.execute(query):
            print(_row)

    print(type(query))
    return (query,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Reading SQL query results into a Pandas data frame

    Here's how we can put the result into a pandas data frame. The selectable and a connection (or engine) are the required arguments for `read_sql`. In this example, we also set the `ID` column to be used as the data frame index.
    """)
    return


@app.cell
def _(engine, pd, query):
    with engine.connect() as _conn:
        df = pd.read_sql(query, _conn, index_col='ID')

    df.head()
    return (df,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Writing a Pandas data frame to a database

    In the next example, we'll use the `shipping_tables.xlsx` table that we used in the [Reading poorly structured Excel files with Pandas notebook](Pandas_messy_data.ipynb).

    First, we'll get the data into data frames.
    """)
    return


@app.cell
def _(pd):
    orders = pd.read_excel('shipping_tables.xlsx', header=1, usecols='B:F')
    orders.head()
    return (orders,)


@app.cell
def _(load_workbook, pd):
    _wb = load_workbook(filename='shipping_tables.xlsx')
    _sheet = _wb['shipping_rates']
    _lookup_table = _sheet.tables['ship_cost']
    _data = _sheet[_lookup_table.ref]

    _rows_list = []
    for _row in _data:
        _cols = []
        for _col in _row:
            _cols.append(_col.value)
        _rows_list.append(_cols)

    ship_rates = pd.DataFrame(data=_rows_list[1:], index=None, columns=_rows_list[0])
    ship_rates.head()
    return (ship_rates,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    Now we have two Pandas data frames and can create a database and add the data.

    In SQLAlchemy 2.x, writes should happen inside a transaction. `engine.begin()` opens a connection and commits automatically when the block finishes (or rolls back if there is an error).
    """)
    return


@app.cell
def _(create_engine, orders, ship_rates):
    # Create the database engine.
    shipping_engine = create_engine('sqlite:///shipping.sqlite')

    # Use the `to_sql` function to send the data frames to the database.
    with shipping_engine.begin() as _conn:
        orders.to_sql('orders', _conn, if_exists='replace')
        ship_rates.to_sql('ship_rates', _conn, if_exists='replace')
    return (shipping_engine,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    Finally, run a raw SQL query to see that the data are there. In SQLAlchemy 2.x, `engine.execute()` has been removed, so we use a connection with `exec_driver_sql()` (or `text()`).
    """)
    return


@app.cell
def _(shipping_engine):
    with shipping_engine.connect() as _conn:
        _rows = _conn.exec_driver_sql("SELECT * FROM orders").fetchall()

    _rows
    return


if __name__ == "__main__":
    app.run()
