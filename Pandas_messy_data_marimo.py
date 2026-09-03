# /// script
# dependencies = ["xlrd"]
# ///

import marimo

__generated_with = "0.24.0"
app = marimo.App()


@app.cell
def _():
    import marimo as mo

    return (mo,)


@app.cell
def _():
    import subprocess

    return (subprocess,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    # Reading poorly structured Excel files with Pandas

    This is a walkthrough of [Chris Moffitt's Practical Business Python](https://pbpython.com/pandas-excel-range.html) blog post from Oct 19, 2020.

    This post came across my feed right as we were looking at Pandas and I think has great examples with real-world datasets. In an ideal world, all data would be nicely formatted and easy to work with...that world does not exist...data are messy and people don't follow best practices in formatting files.

    Note that this tutorial requires openpyxl >= 3.0.4
    """)
    return


@app.cell
def _(subprocess):
    # Download the example file to the current directory.
    #! wget https://github.com/chris1610/pbpython/raw/master/data/shipping_tables.xlsx
    subprocess.call(['wget', 'https://github.com/chris1610/pbpython/raw/master/data/shipping_tables.xlsx'])
    return


@app.cell
def _():
    # Need a module for Excel that isn't installed on HiPerGato.
    #  `pip install MODULE --user` is the command that would install modules
    #  in your user directory.
    # packages added via marimo's package management: xlrd !pip install xlrd --user
    return


@app.cell
def _():
    import pandas as pd
    _df = pd.read_excel('shipping_tables.xlsx')
    # Let's see what happens if we simply try to read this into a dataframe
    _df.head()
    return (pd,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    Simply reading the file in with `pd.read_excel()` gives messy results because the first row of the first sheet has just a date in columns I & J. The real headers are in the second row of the Excel file. Also column A has no data, so we can ignore that. The `pd.read_excel()` function has options to deal with these.
    """)
    return


@app.cell
def _(pd):
    _df = pd.read_excel('shipping_tables.xlsx', header=1, usecols='B:F')
    _df.head()
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    As noted, the `header=1` is a 0-based index to the header row--the second row in this case.

    The `usecols` flag also takes a lot of different formats for the specification, letters, numbers, column names, etc. The original post also looks at using a lambda function to make all the column names lower case so that multiple files with similar column names can be combined.

    Pandas can also read from (and write to) lots of different types of data sources. Check the [I/O section of the Pandas docs](https://pandas.pydata.org/pandas-docs/stable/reference/io.html).
    """)
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Reading from Excel Worksheets, Ranges and Tables

    Also notice above that we got the data from the first worksheet, and nothing with information that the Excel file has two worksheets.

    The example in the file may seem extreem, but again...people are people, publishers are publishers, and there's certainly data out there with these formats. One file may not be an issue to work with by hand, but what if you had hundereds of these files to work with?

    While we can get part of the way there using `sheet_name` in `pd.read_excel` that function doesn't know about Table names:
    """)
    return


@app.cell
def _(pd):
    df_rates = pd.read_excel('shipping_tables.xlsx', sheet_name='shipping_rates')

    df_rates.head()
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    To access named tables within the Excel file, there is another module called `openpyxl`.
    """)
    return


@app.cell
def _():
    from openpyxl import load_workbook

    wb = load_workbook(filename = 'shipping_tables.xlsx') # Notice the different format here where 
                                                          # filename flag is needed
    type(wb) # Like Pandas dataframes, openpyxl adds a data type of workbook.
    return (wb,)


@app.cell
def _(wb):
    wb.sheetnames
    return


@app.cell
def _(wb):
    # Create a sheet variable with the shipping_rates sheet

    sheet = wb['shipping_rates']
    return (sheet,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    Look at the named tables with that sheet.
    """)
    return


@app.cell
def _(sheet):
    sheet.tables.keys()
    return


@app.cell
def _(sheet):
    # Get the Excel range for the ship_cost table:
    # Again, different than the 
    lookup_table = sheet.tables['ship_cost']
    lookup_table.ref
    return (lookup_table,)


@app.cell
def _(lookup_table, pd, sheet):
    # Using the above range, convert the data into a data frame
    data = sheet[lookup_table.ref]
    # Access the data in the table range
    rows_list = []
    for row in data:
        cols = []
    # Loop through each row and get the values in the cells
        for col in row:
            cols.append(col.value)  # Get a list of all columns in each row
        rows_list.append(cols)
    _df = pd.DataFrame(data=rows_list[1:], index=None, columns=rows_list[0])
    # Create a pandas dataframe from the rows_list.
    # The first row is the column names
    _df.head()
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Summary

    Chris also points to a paper that I typically assign later in the semester, [Broman and Woo (2018)](https://www.tandfonline.com/doi/full/10.1080/00031305.2017.1375989), that covers best practices in data organization in spreadsheets.
    """)
    return


if __name__ == "__main__":
    app.run()
