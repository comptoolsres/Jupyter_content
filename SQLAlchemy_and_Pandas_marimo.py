import marimo

__generated_with = "0.24.0"
app = marimo.App()


@app.cell
def _():
    import marimo as mo

    return (mo,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    # Pandas and SQLAlchemy

    Pandas has several functions to work with databases via sqlalchemy.

    This can be an easier way to get data from a database and to pass data into a database.

    Let's start with loading the `world.sqlite` database we've been using.
    """)
    return


@app.cell
def _():
    #!/usr/bin/env python
    from sqlalchemy import create_engine
    from sqlalchemy import MetaData
    from sqlalchemy import Table, Column
    from sqlalchemy import Integer, String
    from sqlalchemy import sql, select, join, desc
    _engine = create_engine('sqlite:////blue/bsc4452/share/Class_Files/data/world.sqlite')
    # Create a Engine object which is our handle into the database.
    conn = _engine.connect()
    # Connect to the database
    # Read the metadata from the existing database.
    #  Since the database already exists and has tables defined, we can create Python objects based on these automatically.
    DBInfo = MetaData(_engine)
    return DBInfo, Table, conn, create_engine, select


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    The cell below is also from the [SQLAlchemy notebook](SQLAlchemy.ipynb) and creates the `city` object with the table metadata, creates the `query` variable. This cell shows how you would execute the query with SQLAlchemy, put the result into the `result` variable and then loops through that printing each row. The last line was added to show the the `query` is of type `<class 'sqlalchemy.sql.selectable.Select'>`, or in words a **SQLAlchemy Selectable**, one of the options for the `pandas.read_sql` function we'll look at next.
    """)
    return


@app.cell
def _(DBInfo, Table, conn, select):
    # Auto-create the city object based on the metadata read into the DBInfo.
    city = Table('city', DBInfo, autoload=True)
    query = select([city.c.ID, city.c.Name, city.c.Population]).limit(10)
    # Mirroring this command: SELECT Id, Name, Population FROM city LIMIT 10;
    result = conn.execute(query)
    for _row in result:
        print(_row)
    print(type(query))
    return (query,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Reading SQL query results into a Pandas data frame

    Here's how we can put the result into a pandas data frame. The Selectable and connection are the required arguments for `read_sql`. In this example, we also add that the ID column should be used for the index of the data frame.
    """)
    return


@app.cell
def _(conn, query):
    import pandas as pd

    df=pd.read_sql(query, conn, index_col='ID')

    df.head()
    return (pd,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    ## Writing a Pandas data frame to a database

    In the next example, we'll use the `shipping_tables.xlsx` table that we used in the [Reading poorly structured Excel files with Pandas notebook](Pandas_messy_data.ipynb).

    First, we'll get the data into the dataframes
    """)
    return


@app.cell
def _(pd):
    orders = pd.read_excel('shipping_tables.xlsx', header=1, usecols='B:F')
    orders.head()
    return (orders,)


@app.cell
def _(pd):
    from openpyxl import load_workbook
    wb = load_workbook(filename='shipping_tables.xlsx')
    sheet = wb['shipping_rates']
    lookup_table = sheet.tables['ship_cost']
    data = sheet[lookup_table.ref]
    rows_list = []
    for _row in data:
        cols = []
        for col in _row:
            cols.append(col.value)
        rows_list.append(cols)
    ship_rates = pd.DataFrame(data=rows_list[1:], index=None, columns=rows_list[0])
    ship_rates.head()
    return (ship_rates,)


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    Now we have two Pandas data frames and can create a database and add the data.
    """)
    return


@app.cell
def _(create_engine, orders, ship_rates):
    _engine = create_engine('sqlite:///shipping.sqlite')
    conn_1 = _engine.connect()
    orders.to_sql('orders', conn_1)
    ship_rates.to_sql('ship_rates', conn_1)
    _engine.execute('SELECT * FROM orders').fetchall()
    return


if __name__ == "__main__":
    app.run()
